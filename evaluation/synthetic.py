"""Deterministic SYNTHETIC resume / job-description generator with latent ground truth.

IMPORTANT: everything produced here is synthetic. Resumes are assembled from templates by this module, and
relevance labels are derived from each persona's *latent* skill set (what the persona "truly" knows), not
from human judgement and not from the system under test. Results on this data validate the pipeline and
expose failure modes; they do NOT estimate accuracy on real-world resumes.
"""
from __future__ import annotations

import io
import random
import textwrap
from dataclasses import dataclass, field

# ------------------------------------------------------------------ vocabulary
FAMILIES = {
    "ml_engineer": ["python", "pytorch", "tensorflow", "scikit-learn", "machine learning", "deep learning", "natural language processing", "computer vision", "docker", "aws", "sql", "mlops", "git", "transformers", "feature engineering"],
    "data_analyst": ["sql", "excel", "power bi", "tableau", "python", "pandas", "statistics", "data visualization", "data analysis", "etl", "r"],
    "backend_dev": ["python", "java", "rest api", "fastapi", "django", "postgresql", "mongodb", "docker", "aws", "microservices", "redis", "git", "ci/cd", "kafka", "unit testing"],
    "frontend_dev": ["javascript", "typescript", "react", "html", "css", "next.js", "git", "unit testing", "figma", "rest api", "graphql", "angular"],
    "devops_engineer": ["docker", "kubernetes", "terraform", "aws", "ci/cd", "linux", "monitoring", "bash", "python", "azure", "git"],
    "data_engineer": ["python", "sql", "apache spark", "airflow", "kafka", "etl", "aws", "data warehousing", "docker", "postgresql", "hadoop", "dbt"],
}
SURFACE = {
    "python": ["Python", "Python3"], "pytorch": ["PyTorch", "Torch"], "tensorflow": ["TensorFlow"], "scikit-learn": ["scikit-learn", "sklearn", "Scikit Learn"],
    "machine learning": ["Machine Learning", "ML"], "deep learning": ["Deep Learning"], "natural language processing": ["NLP", "Natural Language Processing"],
    "computer vision": ["Computer Vision"], "docker": ["Docker"], "aws": ["AWS", "Amazon Web Services"], "sql": ["SQL"], "mlops": ["MLOps", "MLflow"], "git": ["Git"],
    "transformers": ["Transformers", "BERT"], "feature engineering": ["Feature Engineering"], "excel": ["Excel", "Microsoft Excel"], "power bi": ["Power BI"], "tableau": ["Tableau"],
    "pandas": ["Pandas"], "statistics": ["Statistics", "Statistical Analysis"], "data visualization": ["Data Visualization", "Data Visualisation"], "data analysis": ["Data Analysis"],
    "etl": ["ETL", "Data Pipelines"], "r": ["R"], "java": ["Java"], "rest api": ["REST APIs", "RESTful API"], "fastapi": ["FastAPI"], "django": ["Django"], "postgresql": ["PostgreSQL", "Postgres"],
    "mongodb": ["MongoDB", "Mongo"], "microservices": ["Microservices"], "redis": ["Redis"], "ci/cd": ["CI/CD", "Jenkins", "GitHub Actions"], "kafka": ["Kafka", "Apache Kafka"],
    "unit testing": ["Unit Testing", "pytest", "Jest"], "javascript": ["JavaScript", "ES6"], "typescript": ["TypeScript"], "react": ["React", "React.js"], "html": ["HTML5", "HTML"],
    "css": ["CSS3", "CSS"], "next.js": ["Next.js"], "figma": ["Figma"], "graphql": ["GraphQL"], "angular": ["Angular"], "kubernetes": ["Kubernetes", "k8s"], "terraform": ["Terraform"],
    "linux": ["Linux"], "monitoring": ["Prometheus", "Grafana", "Monitoring"], "bash": ["Bash", "Shell Scripting"], "azure": ["Azure"], "apache spark": ["Spark", "PySpark", "Apache Spark"],
    "airflow": ["Airflow", "Apache Airflow"], "data warehousing": ["Snowflake", "Data Warehousing", "BigQuery"], "hadoop": ["Hadoop"], "dbt": ["dbt"],
}
# Descriptions of the same abilities WITHOUT the usual keyword (tests semantic matching).
PARAPHRASE = {
    "pytorch": "Trained and tuned convolutional and recurrent models with a Python tensor library supporting automatic differentiation",
    "tensorflow": "Exported trained networks to a graph format and served them for low-latency inference",
    "scikit-learn": "Built random-forest and gradient-boosted predictors with cross-validated hyperparameter search",
    "machine learning": "Developed predictive models that learn patterns from historical records to forecast customer churn",
    "deep learning": "Trained multi-layer neural architectures for image and text tasks on GPUs",
    "natural language processing": "Created systems that classify, summarise and extract entities from free-form customer messages",
    "computer vision": "Developed image classifiers and detectors that recognise defects on factory-line photographs",
    "docker": "Packaged services into reproducible isolated runtime images shipped to every environment",
    "aws": "Ran workloads on Amazon's public cloud using managed compute, object storage and serverless functions",
    "sql": "Wrote complex relational queries with joins, window functions and aggregations over large tables",
    "mlops": "Automated model training, versioning and release so new models reach production safely",
    "transformers": "Adapted pretrained attention-based language encoders to domain text",
    "feature engineering": "Designed informative input variables from raw logs to improve model accuracy",
    "excel": "Built spreadsheet models with pivot tables and lookup formulas for monthly reporting",
    "power bi": "Published interactive management dashboards on Microsoft's business intelligence platform",
    "tableau": "Designed drill-down visual analytics workbooks for sales leadership",
    "pandas": "Cleaned and reshaped tabular datasets with dataframe libraries for analysis",
    "statistics": "Applied hypothesis tests, confidence intervals and regression to evaluate experiments",
    "data visualization": "Presented findings through charts and visual storytelling for non-technical audiences",
    "data analysis": "Investigated business questions by exploring data, finding trends and recommending actions",
    "etl": "Built scheduled jobs that extract records from source systems, transform them and load a reporting store",
    "java": "Implemented object-oriented server components on the JVM with a dependency-injection framework",
    "rest api": "Designed HTTP endpoints exchanging JSON for web and mobile clients with versioned contracts",
    "fastapi": "Delivered asynchronous Python web services with automatic request validation",
    "django": "Developed full-stack Python web applications with an ORM and admin interface",
    "postgresql": "Modelled and tuned a relational database with indexes, transactions and migrations",
    "mongodb": "Stored semi-structured documents in a document-oriented database with flexible schemas",
    "microservices": "Decomposed a monolith into independently deployable services communicating over queues and HTTP",
    "redis": "Added in-memory key-value caching to cut response times",
    "ci/cd": "Automated build, test and release stages so every merge deploys with no manual steps",
    "kafka": "Processed real-time event streams through a distributed log with consumer groups",
    "unit testing": "Wrote automated tests with fixtures and mocks, running on every commit",
    "javascript": "Built dynamic browser interfaces with the web's scripting language",
    "typescript": "Added static typing to a large browser codebase to catch defects early",
    "react": "Built component-based single-page applications with hooks and client-side state management",
    "html": "Authored semantic, accessible page markup",
    "css": "Styled responsive layouts with flexbox, grid and design tokens",
    "next.js": "Delivered server-rendered React pages with file-based routing",
    "figma": "Created clickable interface prototypes and design systems with shared components",
    "graphql": "Exposed a typed query language over the backend so clients fetch exactly the fields they need",
    "angular": "Developed enterprise single-page applications with dependency injection and RxJS",
    "kubernetes": "Orchestrated container fleets with automated scaling, rolling updates and self-healing",
    "terraform": "Provisioned cloud infrastructure from declarative templates under version control",
    "linux": "Administered servers, managed processes, permissions and system services from the command line",
    "monitoring": "Built metrics dashboards and alerting so on-call engineers detect incidents early",
    "bash": "Automated routine operations with command-line scripts",
    "azure": "Deployed applications on Microsoft's public cloud platform",
    "apache spark": "Processed terabyte-scale datasets in parallel across a distributed compute cluster",
    "airflow": "Orchestrated dependent batch workflows with scheduling, retries and backfills",
    "data warehousing": "Modelled star schemas in a columnar analytical store for reporting",
    "hadoop": "Maintained a distributed file system and batch processing cluster",
    "dbt": "Maintained versioned, tested SQL transformation models in the analytics warehouse",
    "git": "Collaborated through branches, pull requests and code review",
    "r": "Performed statistical modelling and plotting in the statistical computing language",
    "python": "Wrote clean, modular scripting and application code in a dynamic high-level language",
}
VERBS = ["Built", "Developed", "Designed", "Implemented", "Delivered", "Automated", "Improved", "Maintained", "Led", "Created"]
THINGS = {
    "ml_engineer": ["a prediction service", "a training pipeline", "an experiment framework", "a recommendation module", "an evaluation harness"],
    "data_analyst": ["a weekly reporting suite", "a customer segmentation study", "an executive KPI view", "a sales forecasting workbook"],
    "backend_dev": ["an order-management service", "a payments integration", "an authentication module", "a notifications service", "an inventory API"],
    "frontend_dev": ["a customer dashboard", "a checkout flow", "a design-system component library", "an admin console", "a product catalogue"],
    "devops_engineer": ["the release pipeline", "a multi-environment platform", "the incident-response tooling", "a cost-optimisation initiative"],
    "data_engineer": ["a batch ingestion platform", "a streaming ingestion layer", "a data-quality framework", "the analytics warehouse"],
}
FIRST = ["Aarav", "Mei", "Lucas", "Fatima", "Noah", "Sofia", "Kenji", "Amara", "Ivan", "Priya", "Elena", "Omar", "Hannah", "Diego", "Yara", "Tomas", "Zara", "Arjun", "Chloe", "Samuel", "Ananya", "Mateo", "Leila", "Jonas"]
LAST = ["Patel", "Nguyen", "Garcia", "Okafor", "Müller", "Rossi", "Tanaka", "Haddad", "Kowalski", "Singh", "Silva", "Johansson", "Ahmed", "Dubois", "Kim", "Costa", "Novak", "Reyes", "Banerjee", "Fischer", "Mendes", "Larsen", "Iyer", "Brown"]
COMPANIES = ["Northwind Technologies", "BlueRiver Software", "Lumen Analytics Pvt Ltd", "Orbital Systems Inc", "Quartz Labs", "Helix Digital Solutions", "Cascade Consulting Group", "Pioneer Data Services", "Summit Cloud Ltd", "Vertex Infotech", "Aurora Networks", "Nimbus Software Pvt Ltd"]
TITLES = {
    "ml_engineer": ["Machine Learning Engineer", "Data Scientist", "AI Engineer"], "data_analyst": ["Data Analyst", "Business Analyst", "BI Analyst"],
    "backend_dev": ["Backend Engineer", "Software Engineer", "Backend Developer"], "frontend_dev": ["Frontend Developer", "UI Engineer", "Web Developer"],
    "devops_engineer": ["DevOps Engineer", "Site Reliability Engineer", "Cloud Engineer"], "data_engineer": ["Data Engineer", "Big Data Engineer", "Analytics Engineer"],
}
UNIS = ["Riverdale University", "Northfield Institute of Technology", "Lakeshore College", "Westbrook University", "St. Mary's College", "Eastgate Institute of Technology", "Cedar Valley University", "Harborview University"]
DEGREES = [("Bachelor of Technology", "B.Tech", 3), ("Bachelor of Science", "B.Sc", 3), ("Master of Science", "M.Sc", 4), ("Master of Technology", "M.Tech", 4)]
FIELDS = ["Computer Science", "Information Technology", "Statistics", "Electronics Engineering", "Mathematics", "Data Science"]
CERTS = [("AWS Certified Solutions Architect", "Amazon Web Services"), ("Google Professional Data Engineer", "Google Cloud"), ("Microsoft Azure Fundamentals", "Microsoft"), ("Deep Learning Specialization", "Coursera"),
         ("Certified Kubernetes Administrator", "Linux Foundation"), ("Tableau Desktop Specialist", "Tableau"), ("TensorFlow Developer Certificate", "Google"), ("Meta Front-End Developer Certificate", "Coursera")]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
TODAY = (2025, 6)
LAYOUTS = ["classic", "pipes", "docx", "minimal"]


@dataclass
class Persona:
    pid: str
    family: str
    kind: str  # strong | paraphrased | moderate | weak | stuffed | adjacent
    layout: str
    name: str
    email: str
    phone: str
    linkedin: str
    github: str
    true_skills: set
    years: float
    named_skills: set  # skills literally named in the resume text
    jobs: list = field(default_factory=list)  # (title, org, start(y,m), end(y,m)|None)
    education: list = field(default_factory=list)  # (degree_label, field, institution, year_start, year_end, level)
    projects: list = field(default_factory=list)  # (name, [skills], text)
    certs: list = field(default_factory=list)
    sections: set = field(default_factory=set)
    structure: list = field(default_factory=list)  # [(kind, text)]
    file: str = ""

    def truth(self) -> dict:
        return {
            "name": self.name, "email": self.email, "phone": "".join(c for c in self.phone if c.isdigit()),
            "linkedin": self.linkedin, "github": self.github,
            "skills": sorted(self.named_skills),
            "edu_institution": [e[2] for e in self.education], "edu_level": [e[5] for e in self.education],
            "exp_title": [j[0] for j in self.jobs], "exp_org": [j[1] for j in self.jobs],
            "exp_dates": [f"{j[2][0]}-{j[2][1]:02d}>{'present' if j[3] is None else f'{j[3][0]}-{j[3][1]:02d}'}" for j in self.jobs],
            "project_name": [p[0] for p in self.projects], "cert_name": [c[0] for c in self.certs],
        }


def _fmt(d, style):
    y, m = d
    return f"{MONTHS[m - 1]} {y}" if style != "minimal" else f"{m:02d}/{y}"


def _skill_bullets(rng, persona_family, skills, named: bool):
    out = []
    for s in skills:
        if named:
            surf = rng.choice(SURFACE.get(s, [s.title()]))
            out.append(f"{rng.choice(VERBS)} {rng.choice(THINGS[persona_family])} using {surf}" + rng.choice(["", " and automated testing", " to support business teams", " for production use"]) + ".")
        else:
            out.append(PARAPHRASE.get(s, f"Worked on tasks involving {s}") + ".")
    return out


def make_persona(rng: random.Random, idx: int, family: str, kind: str, layout: str) -> Persona:
    core = FAMILIES[family]
    k = len(core)
    n_true = {"strong": rng.randint(round(.8 * k), k), "paraphrased": rng.randint(round(.8 * k), k), "moderate": rng.randint(round(.45 * k), round(.65 * k)),
              "weak": rng.randint(1, max(2, round(.2 * k))), "stuffed": rng.randint(1, 3), "adjacent": rng.randint(round(.3 * k), round(.5 * k))}[kind]
    true = set(rng.sample(core, n_true))
    if kind == "adjacent":  # drawn from another family's pool partly
        other = rng.choice([f for f in FAMILIES if f != family])
        true = set(rng.sample(FAMILIES[other], min(n_true, len(FAMILIES[other]))))
    years = {"strong": rng.choice([3, 4, 5, 6, 8]), "paraphrased": rng.choice([3, 4, 5, 6]), "moderate": rng.choice([1, 2, 3, 4]), "weak": rng.choice([0, 1, 2]), "stuffed": rng.choice([0, 1]), "adjacent": rng.choice([2, 3, 4])}[kind]
    name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
    slug = "".join(c for c in name.lower() if c.isalpha() or c == " ").replace(" ", ".")
    p = Persona(pid=f"R{idx:03d}", family=family, kind=kind, layout=layout, name=name, email=f"{slug}{idx}@mail-example.com",
                phone=f"+1 {rng.randint(200, 999)}-{rng.randint(200, 999)}-{rng.randint(1000, 9999)}", linkedin=f"linkedin.com/in/{slug.replace('.', '-')}-{idx}",
                github=f"github.com/{slug.replace('.', '')}{idx}", true_skills=set(true), years=years, named_skills=set())
    p.sections = {"experience", "education", "skills"}
    for opt, pr in (("summary", .7), ("projects", .75), ("certifications", .45)):
        if rng.random() < pr:
            p.sections.add(opt)
    titles = TITLES[family if kind != "adjacent" else other] if kind != "adjacent" else TITLES[other]
    # experience
    ey, em = TODAY
    n_jobs = 0 if years == 0 else (1 if years <= 2 else 2 if years <= 5 else 3)
    cur_end = None
    start_y = ey - int(years)
    spans = []
    if n_jobs:
        per = max(1, int(years * 12 / n_jobs))
        end_idx = ey * 12 + em - 1
        for j in range(n_jobs):
            s_idx = end_idx - per
            spans.append(((s_idx // 12, s_idx % 12 + 1), None if j == 0 else (end_idx // 12, end_idx % 12 + 1)))
            end_idx = s_idx - rng.randint(1, 2)
    comps = rng.sample(COMPANIES, max(1, n_jobs))
    skills_pool = sorted(true)
    rng.shuffle(skills_pool)
    for j in range(n_jobs):
        p.jobs.append((rng.choice(titles), comps[j], spans[j][0], spans[j][1]))
    # which true skills are literally named
    para = kind == "paraphrased"
    named_pool = set() if para else set(true)
    p.named_skills = set(named_pool)
    # education
    deg = rng.choice(DEGREES)
    ys = ey - int(years) - 4 - rng.randint(0, 1)
    p.education.append((deg[0] if layout != "pipes" else deg[1], rng.choice(FIELDS), rng.choice(UNIS), ys, ys + (4 if deg[2] == 3 else 2), deg[2]))
    # projects
    if "projects" in p.sections:
        for t in range(rng.randint(1, 3)):
            ps = rng.sample(skills_pool, min(len(skills_pool), rng.randint(1, 3)))
            nm = rng.choice(["Smart", "Open", "Rapid", "Unified", "Insight", "Nova", "Atlas", "Pulse"]) + " " + rng.choice(["Tracker", "Pipeline", "Studio", "Monitor", "Planner", "Engine"]) + f" {t + 1}"
            p.projects.append((nm, ps, ""))
    if "certifications" in p.sections:
        p.certs = rng.sample(CERTS, rng.randint(1, 2))
    # --------------------------------------------------------------- structure
    st = p.structure
    classic = layout in ("classic", "docx")
    H = (lambda t: t.upper()) if layout == "classic" else (lambda t: t.title() if layout != "minimal" else t)
    heads = {"experience": "Work Experience" if layout != "minimal" else "Professional Background", "education": "Education" if layout != "minimal" else "Academic Background",
             "skills": "Technical Skills" if layout != "minimal" else "Technical Proficiency", "projects": "Projects" if layout != "minimal" else "Selected Projects",
             "certifications": "Certifications", "summary": "Summary"}
    hd = lambda k: ("h", H(heads[k]) + (":" if layout == "pipes" else ""))
    bullet = "bullet" if layout != "minimal" else "line"
    st.append(("name", p.name))
    st.append(("contact", f"{p.email} | {p.phone} | {p.linkedin} | {p.github}"))
    if "summary" in p.sections:
        st += [hd("summary"), ("line", f"{rng.choice(titles)} with {int(years)} years of experience. " + ("Skilled in " + ", ".join(rng.choice(SURFACE.get(s, [s])) for s in skills_pool[:3]) + "." if not para else "Focused on building reliable, maintainable solutions."))]
    # skills section text
    if kind == "stuffed":
        stuff = rng.sample(core, min(len(core), rng.randint(round(.8 * k), k)))
        listed = stuff
        p.named_skills |= set(stuff)
    else:
        listed = [s for s in skills_pool if s in named_pool]
    # experience
    st.append(hd("experience"))
    used = skills_pool[:]
    for ji, (ti, org, s, e) in enumerate(p.jobs):
        dr = f"{_fmt(s, layout)} - {_fmt(e, layout) if e else 'Present'}"
        if layout == "classic" or layout == "docx":
            st += [("line", f"{ti} | {org} | {dr}")]
        elif layout == "pipes":
            st += [("line", f"{org}"), ("line", f"{ti}    {dr}")]
        else:
            st += [("line", f"{ti}, {org} ({dr})")]
        chunk = used[ji::max(1, n_jobs)]
        chunk = chunk[:5] or used[:2]
        for ln in _skill_bullets(rng, p.family if kind != "adjacent" else other, chunk, not para):
            st.append((bullet, ln))
        st.append(("blank", ""))
    if n_jobs == 0:
        st.append(("line", "Fresher: coursework and academic projects listed below."))
    # education
    st.append(hd("education"))
    for (dg, fl, un, y0, y1, lv) in p.education:
        if layout == "classic" or layout == "docx":
            st.append(("line", f"{dg} in {fl}, {un}, {y0} - {y1}"))
        elif layout == "pipes":
            st += [("line", f"{un}"), ("line", f"{dg} ({fl})    {y0} - {y1}")]
        else:
            st.append(("line", f"{un} - {dg} in {fl} ({y0}-{y1})"))
    if "projects" in p.sections:
        st.append(hd("projects"))
        for (nm, ps, _) in p.projects:
            if para:
                st += [("line", nm)] + [(bullet, PARAPHRASE[s] + ".") for s in ps]
            else:
                st += [("line", f"{nm} | " + ", ".join(rng.choice(SURFACE.get(s, [s])) for s in ps))]
                st += [(bullet, f"{rng.choice(VERBS)} {rng.choice(THINGS[p.family])} with {rng.choice(SURFACE.get(ps[0], [ps[0]]))}.")]
                p.named_skills |= set(ps)
            st.append(("blank", ""))
    if "certifications" in p.sections:
        st.append(hd("certifications"))
        for (cn, iss) in p.certs:
            st.append((bullet, f"{cn} - {iss}, {rng.randint(2019, 2024)}"))
    if listed:
        st.append(hd("skills"))
        shown = [rng.choice(SURFACE.get(s, [s.title()])) for s in listed]
        if kind == "stuffed":
            rng.shuffle(shown)
        st.append(("line", "Skills: " + ", ".join(shown)))
    # recompute named skills accurately: project/exp bullets name them
    if not para:
        p.named_skills = set(p.true_skills) | (set(stuff) if kind == "stuffed" else set())
    else:
        p.named_skills = set()
    return p


# ------------------------------------------------------------------ renderers
def unicode_fonts():
    """(regular, bold) TTF paths from matplotlib's bundled DejaVu fonts, or None if unavailable."""
    try:
        import matplotlib
        from pathlib import Path

        d = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        if (d / "DejaVuSans.ttf").exists():
            return str(d / "DejaVuSans.ttf"), str(d / "DejaVuSans-Bold.ttf")
    except Exception:
        pass
    return None


def render_pdf(structure, unicode_font=False) -> bytes:
    import fitz

    doc = fitz.open()
    page, y = None, 0
    uf = unicode_fonts() if unicode_font else None

    def new():
        nonlocal page, y
        page = doc.new_page(width=595, height=842)
        y = 56
        if uf:
            page.insert_font(fontname="uni", fontfile=uf[0])
            page.insert_font(fontname="unib", fontfile=uf[1])
    new()
    for kind, text in structure:
        size = {"name": 18, "h": 11.5}.get(kind, 10)
        font = ("unib" if uf else "hebo") if kind in ("name", "h") else ("uni" if uf else "helv")
        if kind == "blank":
            y += 7
            continue
        if kind == "h":
            y += 8
        prefix = "• " if kind == "bullet" else ""
        wrap = textwrap.wrap(prefix + text, (78 if uf else 92) if kind != "name" else 50, subsequent_indent="  " if prefix else "") or [""]
        for ln in wrap:
            if y > 790:
                new()
            page.insert_text((50, y), ln, fontname=font, fontsize=size)
            y += size + 3.4
    if uf:
        try:
            doc.subset_fonts()  # keep generated PDFs small
        except Exception:
            pass
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def render_docx(structure) -> bytes:
    import docx

    d = docx.Document()
    for kind, text in structure:
        if kind == "name":
            d.add_heading(text, level=0)
        elif kind == "h":
            d.add_heading(text, level=1)
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "blank":
            continue
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_scanned_pdf(structure, dpi=200, jpeg=False) -> bytes:
    """Rasterise a text PDF into an image-only PDF (no text layer) to exercise OCR."""
    import fitz

    src = fitz.open(stream=render_pdf(structure), filetype="pdf")
    out = fitz.open()
    for pg in src:
        pix = pg.get_pixmap(dpi=dpi)
        np_ = out.new_page(width=pg.rect.width, height=pg.rect.height)
        np_.insert_image(np_.rect, stream=pix.tobytes("jpeg", jpg_quality=70) if jpeg else pix.tobytes("png"))
    data = out.tobytes(garbage=3, deflate=True)
    return data


def render(p: Persona, scanned=False) -> tuple[str, bytes]:
    if scanned:
        return f"{p.pid}_scan.pdf", render_scanned_pdf(p.structure)
    if p.layout == "docx":
        return f"{p.pid}.docx", render_docx(p.structure)
    return f"{p.pid}.pdf", render_pdf(p.structure)


# ------------------------------------------------------------------ jobs and labels
JOBS = [
    {"id": "J1", "family": "ml_engineer", "required": ["python", "machine learning", "deep learning", "pytorch", "sql"], "preferred": ["docker", "aws", "natural language processing"], "min_years": 2,
     "text": """Machine Learning Engineer
About us: a fast-growing product company.
Responsibilities:
- Train, evaluate and deploy machine learning models
- Collaborate with product teams
Requirements:
- 2+ years of experience in applied machine learning
- Strong Python and SQL skills
- Experience with deep learning frameworks such as PyTorch
- Bachelor's degree in Computer Science or related field
Nice to have:
- Docker, AWS
- Experience with NLP
"""},
    {"id": "J2", "family": "data_analyst", "required": ["sql", "excel", "data visualization", "statistics"], "preferred": ["python", "tableau", "power bi"], "min_years": 1,
     "text": """Data Analyst
Requirements:
- 1+ years of experience in data analysis
- Proficiency in SQL and Excel
- Experience building dashboards and data visualization with Tableau or Power BI
- Solid understanding of statistics
Preferred:
- Python or R for analysis
"""},
    {"id": "J3", "family": "backend_dev", "required": ["python", "rest api", "postgresql", "docker", "git"], "preferred": ["aws", "kafka", "ci/cd"], "min_years": 3,
     "text": """Backend Engineer
Responsibilities:
- Design and build scalable REST APIs
- Own services end to end
Requirements:
- 3+ years of backend development experience
- Python with FastAPI or Django
- PostgreSQL, Docker and Git
- Experience designing RESTful APIs
Bonus points:
- AWS, Kafka, CI/CD pipelines
"""},
    {"id": "J4", "family": "frontend_dev", "required": ["javascript", "react", "html", "css", "typescript"], "preferred": ["next.js", "unit testing", "graphql"], "min_years": 2,
     "text": """Frontend Developer
Must have:
- 2+ years building web applications
- Expert JavaScript and TypeScript
- React, HTML and CSS
Nice to have:
- Next.js, GraphQL, automated unit testing
"""},
    {"id": "J5", "family": "devops_engineer", "required": ["docker", "kubernetes", "terraform", "ci/cd", "linux"], "preferred": ["aws", "monitoring", "bash"], "min_years": 3,
     "text": """DevOps Engineer
Requirements:
- 3+ years in DevOps or SRE roles
- Kubernetes and Docker in production
- Infrastructure as code with Terraform
- Build CI/CD pipelines
- Strong Linux skills
Preferred:
- AWS or Azure certifications; observability tooling like Prometheus
"""},
    {"id": "J6", "family": "data_engineer", "required": ["python", "sql", "apache spark", "airflow", "etl"], "preferred": ["kafka", "aws", "data warehousing"], "min_years": 3,
     "text": """Data Engineer
Responsibilities:
- Build and maintain data pipelines
Requirements:
- 3+ years of experience as a data engineer
- Python and advanced SQL
- Apache Spark and Airflow
- Experience designing ETL processes
Nice to have:
- Kafka, AWS, a cloud data warehouse
"""},
    {"id": "J7", "family": "ml_engineer", "required": ["natural language processing", "transformers", "python", "deep learning"], "preferred": ["pytorch", "mlops"], "min_years": 2,
     "text": """NLP Engineer
We are looking for someone to build language understanding systems.
- Experience with natural language processing and transformer models
- Python, deep learning
- Plus: PyTorch, MLOps
"""},
    {"id": "J8", "family": "backend_dev", "required": ["java", "microservices", "rest api", "postgresql", "unit testing"], "preferred": ["kafka", "docker", "redis"], "min_years": 4,
     "text": """Senior Java Engineer
Requirements:
- 4+ years of Java development
- Microservices architecture and RESTful APIs
- PostgreSQL
- Unit testing discipline
Preferred:
- Kafka, Docker, Redis
"""},
]


def grade(p: Persona, job: dict) -> int:
    """Graded relevance 0-3 from latent truth only (see module docstring)."""
    req, pref = set(job["required"]), set(job["preferred"])
    cov = len(req & p.true_skills) / len(req)
    pcov = len(pref & p.true_skills) / len(pref) if pref else 0
    g = 3 if (cov >= .75 and pcov >= .25) or cov >= .9 else 2 if cov >= .5 else 1 if cov >= .25 else 0
    if p.years < job["min_years"] - 2 and g > 0:
        g -= 1
    return g


def build_corpus(seed=7, per_family=None):
    """Return (personas, jobs). Pool design per family: 3 strong, 2 paraphrased, 2 moderate, 2 weak, 1 stuffed + 6 adjacent overall."""
    rng = random.Random(seed)
    personas, idx = [], 0
    mix = per_family or [("strong", 3), ("paraphrased", 2), ("moderate", 2), ("weak", 2), ("stuffed", 1), ("adjacent", 1)]
    for fam in FAMILIES:
        for kind, n in mix:
            for _ in range(n):
                idx += 1
                personas.append(make_persona(rng, idx, fam, kind, LAYOUTS[idx % len(LAYOUTS)]))
    return personas, JOBS
